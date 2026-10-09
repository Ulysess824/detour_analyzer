"""Excel export of the comparison: two tables, "faltante" for the internal model, a plain-words reason for every excluded row."""

import numpy as np
import openpyxl
import pandas as pd

from src.utils.excel_utils import COMPARISON_COLUMNS, EXCLUDED_COLUMNS, comparison_sheet, excluded_sheet, write_workbook
from src.utils.planner_utils import LOW_HISTORY, MISSING_INTERNAL, SUBGRADE


def test_comparison_sheet_has_only_the_keys_and_the_four_values_and_marks_the_missing_internal():
    table = pd.DataFrame(
        {"mes": "2026-09", "planta": ["B", "A"], "sku": ["K/01/110gsm/2100mm", "K/01/125gsm/2100mm"], "strategy": "VMI", "real": [10.04, 20.0], "planner": [11.0, 19.0], "internal": [np.nan, 18.26], "ml_mean": [9.0, 21.0], "history": [30, 30]}
    )
    out = comparison_sheet(table)
    assert list(out.columns) == COMPARISON_COLUMNS
    assert out["Planta"].tolist() == ["A", "B"]  # sorted by planta
    assert out["Modelo interno (TO)"].tolist() == [18.3, "faltante"]
    assert out["Real (TO)"].tolist() == [20.0, 10.0]


def test_every_excluded_row_gets_a_reason_and_an_explanation():
    excluded = pd.DataFrame(
        {
            "planta": "A", "mes": "2026-09", "sku": ["S1", "S2", "S3", "S4"], "strategy": "VMI", "planner": [1.0, 2.0, 3.0, 4.0],
            "reason": [SUBGRADE, LOW_HISTORY, MISSING_INTERNAL, MISSING_INTERNAL],
            "detail": ["en los datos: KS/257/215gsm/1800mm", "2 meses antes del origen", "0 en todos los meses", "no está en el archivo del modelo interno"],
        }
    )
    out = excluded_sheet(excluded)
    assert list(out.columns) == EXCLUDED_COLUMNS
    by = out.set_index("SKU")
    assert by.loc["S1", "Motivo"] == "Subgrado distinto" and "KS/257/215gsm/1800mm" in by.loc["S1", "Explicación"]
    assert "solo 2 meses de consumo" in by.loc["S2", "Explicación"]
    assert by.loc["S3", "Motivo"] == by.loc["S4", "Motivo"] == "Modelo interno sin pronóstico"
    assert "0 en todos los meses" in by.loc["S3", "Explicación"] and "no está en el archivo" in by.loc["S4", "Explicación"]
    assert (out["Explicación"].str.len() > 20).all()


def test_workbook_has_two_sheets_each_with_an_excel_table(tmp_path):
    sheets = {"Comparación": pd.DataFrame({"SKU": ["a"], "Real (TO)": [1.5], "Modelo interno (TO)": ["faltante"]}), "Excluidos": pd.DataFrame({"SKU": ["b"], "Motivo": ["x"]})}
    path = tmp_path / "c.xlsx"
    write_workbook(path, sheets)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Comparación", "Excluidos"]
    assert all(len(ws.tables) == 1 for ws in wb)
    assert wb["Comparación"]["C2"].value == "faltante" and wb["Comparación"]["B2"].value == 1.5
