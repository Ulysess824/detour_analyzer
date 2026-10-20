"""The example of the raw export shown in Datos: the table on the screen and the file to download are the same rows, and the loader reads them."""

import datetime
import io

import openpyxl
from fastapi.testclient import TestClient

from src.app.api import create_app
from src.app.example_service import EXAMPLE_NAME, ROWS, example_workbook
from src.app.paths import AppPaths
from src.engine.transform_utils import read_export


def test_the_table_on_the_screen_and_the_file_to_download_are_the_same(tmp_path):
    client = TestClient(create_app(AppPaths(tmp_path)))
    table = client.get("/api/datos/formato").json()
    assert [c["letra"] for c in table["columnas"]] == list("ABCDEF")
    assert [c["titulo"] for c in table["columnas"]] == ["Customer", "", "Material", "", "Day", "TO"]  # B and D have no title in a real export

    response = client.get("/api/descargas/ejemplo-exporte")
    assert response.status_code == 200 and f'filename="{EXAMPLE_NAME}"' in response.headers["content-disposition"]
    sheet = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert [list(row) for row in sheet.iter_rows(min_row=2, values_only=True)] == table["filas"]
    assert list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))) == ["Customer", None, "Material", None, "Day", "TO"]


def test_the_loader_reads_the_example_it_gives(tmp_path):
    path = tmp_path / EXAMPLE_NAME
    path.write_bytes(example_workbook()[1])
    records = read_export(path, drop_zeros=True)
    assert len(records) == len(ROWS) and {r[0] for r in records} == {"DEMO1", "DEMO2"}
    assert records[0] == ("DEMO1", "K/01/125gsm/2000mm/1200-1400", "2026-03-02", 12.345)  # sku = the description, day as yyyy-mm-dd
    assert all(r[3] > 0 for r in records)


def test_only_a_day_written_as_text_dd_mm_yyyy_is_read(tmp_path):
    r"""What the screen warns about: an Excel date or another text format is dropped without any message."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Customer", None, "Material", None, "Day", "TO"])
    for day in ("17.02.2026", datetime.datetime(2026, 2, 18), "2026-02-19"):
        sheet.append(["DEMO1", "CLIENTE", "900001", "K/01/125gsm/2000mm/1200-1400", day, 5.0])
    path = tmp_path / "days.xlsx"
    book.save(path)
    assert [r[2] for r in read_export(path)] == ["2026-02-17"]
