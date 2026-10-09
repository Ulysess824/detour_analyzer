"""Internal model workbook (every planta) and its join with the planner comparison: totals check, missing forecasts, the planta in the key."""

import numpy as np
import openpyxl
import pandas as pd
import pytest

from src.utils.internal_io_utils import OUTPUT_COLUMNS, mark_missing, parse_month_column, planta_is_all_zero, read_internal_workbook, replace_planta
from src.utils.planner_utils import MISSING_INTERNAL, add_internal

HEADER = ["Customer", "Grade", "Sub-grade", "Grammage", "Width", "Product", "06.2026", "07.2026"]


def write_workbook(path, rows, totals=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["", "", "", "", "", "", "Proposed Forecast", "Proposed Forecast"])
    ws.append(HEADER)
    for row in rows:
        ws.append(row)
    if totals:
        ws.append(["Result", None, None, None, None, None, sum(r[6] for r in rows), sum(r[7] for r in rows)])
        ws.append([None, None, None, None, None, None, sum(r[6] for r in rows), sum(r[7] for r in rows)])
    wb.save(path)


ROWS = [["SCAN", "K", 1, 110, 2100, 111, 0, 0], ["PCEL", "HP2", 1, 115, 2190, 222, 101, 109], ["PCEL", "KS", 257, 215, 2450, 333, 0, 0], ["SALI", "TSL", 1, 100, 2450, 444, 0, 12]]


def test_month_columns():
    assert parse_month_column("06.2026") == "2026-06"
    assert parse_month_column("9.2026") == "2026-09"
    with pytest.raises(ValueError):
        parse_month_column("13.2026")


def test_workbook_is_read_without_the_result_and_empty_rows(tmp_path):
    path = tmp_path / "i.xlsx"
    write_workbook(path, ROWS)
    long = read_internal_workbook(path)
    assert list(long.columns) == OUTPUT_COLUMNS
    assert len(long) == 4 * 2  # four SKUs, two months; no "Result" row
    assert set(long["sku_planner"]) == {"K/01/110gsm/2100mm", "HP2/01/115gsm/2190mm", "KS/257/215gsm/2450mm", "TSL/01/100gsm/2450mm"}
    assert long.loc[(long["sku_planner"] == "HP2/01/115gsm/2190mm") & (long["mes"] == "2026-07"), "forecast_to"].item() == 109.0


def test_a_result_row_that_is_not_the_sum_is_rejected(tmp_path):
    path = tmp_path / "bad.xlsx"
    write_workbook(path, ROWS)
    wb = openpyxl.load_workbook(path)
    wb.active.cell(row=7, column=7).value = 999  # the Result row of June
    wb.save(path)
    with pytest.raises(ValueError, match="Result"):
        read_internal_workbook(path)


def test_a_planta_of_zeros_is_replaced_by_the_earlier_values(tmp_path):
    path = tmp_path / "i.xlsx"
    write_workbook(path, ROWS)
    long = read_internal_workbook(path)
    assert planta_is_all_zero(long, "SCAN") and not planta_is_all_zero(long, "PCEL")
    earlier = pd.DataFrame({"planta": ["SCAN", "SCAN"], "mes": ["2026-06", "2026-07"], "sku_planner": ["K/01/110gsm/2100mm"] * 2, "product_sap": [111, 111], "forecast_to": [40.0, 41.0]})
    out = replace_planta(long, earlier, "SCAN")
    scan = out[out["planta"] == "SCAN"]
    assert scan["forecast_to"].tolist() == [40.0, 41.0]
    assert len(out[out["planta"] != "SCAN"]) == 6  # the other plantas are untouched


def test_zero_in_every_month_is_missing_and_a_partial_zero_is_kept(tmp_path):
    path = tmp_path / "i.xlsx"
    write_workbook(path, ROWS)
    marked = mark_missing(read_internal_workbook(path))
    kept = marked.set_index(["sku_planner", "mes"])["forecast_to"]
    assert np.isnan(kept[("KS/257/215gsm/2450mm", "2026-06")]) and np.isnan(kept[("KS/257/215gsm/2450mm", "2026-07")])
    assert kept[("TSL/01/100gsm/2450mm", "2026-06")] == 0.0  # 0 in June but 12 in July: a real forecast of 0
    assert kept[("TSL/01/100gsm/2450mm", "2026-07")] == 12.0


def test_internal_is_joined_by_planta_and_missing_skus_are_flagged():
    table = pd.DataFrame({"planta": ["A", "B", "A", "A"], "mes": "2026-09", "sku": ["K/01/110gsm/2100mm", "K/01/110gsm/2100mm", "K/01/125gsm/2100mm", "K/01/160gsm/2100mm"], "strategy": "VMI", "planner": [1.0, 2.0, 3.0, 4.0]})
    internal = pd.DataFrame(
        {
            "planta": ["A", "B", "A", "A"],
            "mes": "2026-09",
            "sku_planner": ["K/01/110gsm/2100mm", "K/01/110gsm/2100mm", "K/01/125gsm/2100mm", "K/01/110gsm/2100mm"],
            "product_sap": 1,
            "forecast_to": [10.0, 20.0, np.nan, 99.0],  # the last row is the same SKU in another month
        }
    )
    internal.loc[3, "mes"] = "2026-08"
    out, missing = add_internal(table, internal, "2026-09")
    assert len(out) == 4  # a SKU in two plantas or two months does not duplicate rows
    by = out.set_index(["planta", "sku"])
    assert by.loc[("A", "K/01/110gsm/2100mm"), "internal"] == 10.0 and by.loc[("B", "K/01/110gsm/2100mm"), "internal"] == 20.0
    assert by["internal_status"].to_dict() == {("A", "K/01/110gsm/2100mm"): "ok", ("B", "K/01/110gsm/2100mm"): "ok", ("A", "K/01/125gsm/2100mm"): "faltante", ("A", "K/01/160gsm/2100mm"): "faltante"}
    reasons = dict(zip(missing["sku"], missing["detail"]))
    assert set(missing["reason"]) == {MISSING_INTERNAL}
    assert reasons == {"K/01/125gsm/2100mm": "0 en todos los meses", "K/01/160gsm/2100mm": "no está en el archivo del modelo interno"}
