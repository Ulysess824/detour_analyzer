"""Loading new SAP exports from the app: review against the master, conflicts, checks, backup, undo, and the API around it."""

import io
import os
import time
from datetime import date

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from src.app import data_service
from src.app.api import create_app
from src.app.paths import AppPaths
from src.engine.transform_utils import write_csv
from tests.conftest import make_consumption
from tests.test_app_forecast import wait

TODAY = date(2026, 10, 20)


def xlsx(df: pd.DataFrame, header=("Customer", "Day", "Material", None, "TO")) -> bytes:
    r"""A long SAP export: planta, day, material code, material description, consumption."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(list(header))
    for r in df.itertuples():
        sheet.append([r.planta, r.fecha.strftime("%d.%m.%Y"), "1000", r.sku, float(r.consumo)])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def table(df: pd.DataFrame) -> list[tuple]:
    return list(zip(df["planta"], df["sku"], df["fecha"].dt.strftime("%Y-%m-%d"), df["consumo"]))


@pytest.fixture(scope="module")
def toy():
    return make_consumption()


@pytest.fixture()
def paths(tmp_path, toy):
    r"""A folder whose master has the toy data to 2025-05-31."""
    write_csv(table(toy[toy["fecha"] <= "2025-05-31"]), AppPaths(tmp_path).consumption_csv)
    return AppPaths(tmp_path)


def load(paths, df, name="export.xlsx", policy="detener"):
    staged = data_service.stage(paths, [(name, xlsx(df))])
    return staged["id"], data_service.review(paths, staged["id"], policy)


def read_master(paths) -> pd.DataFrame:
    return pd.read_csv(paths.consumption_csv, parse_dates=["fecha"], float_precision="round_trip")


# ---- staging -----------------------------------------------------------------------------------------------------------------


def test_only_excel_files_are_accepted_and_names_are_cleaned(paths, toy):
    with pytest.raises(ValueError, match="Solo se aceptan"):
        data_service.stage(paths, [("datos.csv", b"x")])
    with pytest.raises(ValueError, match="ningún archivo"):
        data_service.stage(paths, [])
    staged = data_service.stage(paths, [("../../evil.xlsx", xlsx(toy.head(5)))])
    saved = list((paths.staging / staged["id"]).iterdir())
    assert [f.name for f in saved] == ["01_evil.xlsx"] and staged["archivos"][0]["nombre"] == "evil.xlsx"


def test_old_uploads_nobody_finished_are_swept(paths, toy):
    old = data_service.stage(paths, [("a.xlsx", xlsx(toy.head(5)))])["id"]
    past = time.time() - (data_service.STALE_HOURS + 1) * 3600
    os.utime(paths.staging / old, (past, past))
    data_service.stage(paths, [("b.xlsx", xlsx(toy.head(5)))])
    assert not (paths.staging / old).exists()


def test_cancel_removes_the_upload_and_an_unknown_one_is_refused(paths, toy):
    staged = data_service.stage(paths, [("a.xlsx", xlsx(toy.head(5)))])["id"]
    data_service.cancel(paths, staged)
    assert not (paths.staging / staged).exists()
    for bad in (staged, "zzzz", "../x"):
        with pytest.raises(LookupError):
            data_service.review(paths, bad)


def test_a_file_that_is_not_an_export_is_explained(paths):
    staged = data_service.stage(paths, [("roto.xlsx", b"no es un excel")])["id"]
    with pytest.raises(ValueError, match="No se pudo leer roto.xlsx"):
        data_service.review(paths, staged)
    blank = openpyxl.Workbook()
    buffer = io.BytesIO()
    blank.save(buffer)
    staged = data_service.stage(paths, [("en_blanco.xlsx", buffer.getvalue())])["id"]
    with pytest.raises(ValueError, match="No se pudo leer en_blanco.xlsx"):
        data_service.review(paths, staged)
    staged = data_service.stage(paths, [("sin_filas.xlsx", xlsx(make_consumption().head(0)))])["id"]
    with pytest.raises(ValueError, match="no trae filas"):
        data_service.review(paths, staged)
    with pytest.raises(ValueError, match="Política"):
        data_service.review(paths, staged, "otra")


# ---- review and apply --------------------------------------------------------------------------------------------------------


def test_a_new_month_is_added_the_overlap_is_kept_once_and_the_master_is_backed_up(paths, toy):
    before = paths.consumption_csv.read_bytes()
    export = toy[(toy["fecha"] >= "2025-05-15") & (toy["fecha"] <= "2025-06-30")]
    staged, review = load(paths, export)
    summary = review["resumen"]
    overlap = int(((export["fecha"] <= "2025-05-31")).sum())
    assert review["puede_aplicar"] and summary["conflictos"] == 0 and summary["meses_nuevos"] == ["2025-06"]
    assert summary["filas_repetidas_iguales"] == overlap and summary["filas_nuevas"] == len(export) - overlap
    assert summary["ultimo_dia_actual"] == "2025-05-31" and summary["ultimo_dia_resultado"] == "2025-06-30"
    assert paths.consumption_csv.read_bytes() == before  # a review writes nothing to the master

    done = data_service.apply(paths, staged)
    expected = toy[toy["fecha"] <= "2025-06-30"].sort_values(["planta", "sku", "fecha"]).reset_index(drop=True)
    pd.testing.assert_frame_equal(read_master(paths), expected, check_dtype=False)
    assert done["meses_nuevos"] == ["2025-06"] and done["respaldo"] and (paths.backups / done["respaldo"]).read_bytes() == before
    assert not (paths.staging / staged).exists()

    assert data_service.backups(paths)[0]["archivo"] == done["respaldo"]
    assert data_service.undo(paths) == {"restaurado": done["respaldo"]}
    assert paths.consumption_csv.read_bytes() == before and data_service.backups(paths) == []
    with pytest.raises(LookupError):
        data_service.undo(paths)


def test_different_values_for_loaded_days_stop_the_load_unless_the_user_chooses_the_new_ones(paths, toy):
    export = toy[(toy["fecha"] >= "2025-05-20") & (toy["fecha"] <= "2025-06-30")].copy()
    victim = export[export["fecha"] <= "2025-05-31"].iloc[0]
    export.loc[victim.name, "consumo"] = victim["consumo"] + 7.0

    staged, stop = load(paths, export)
    assert not stop["puede_aplicar"] and stop["resumen"]["conflictos"] == 1
    conflict = stop["hallazgos"][0]
    assert conflict["codigo"] == "conflictos" and conflict["nivel"] == "error" and conflict["ejemplos"][0]["nuevo"] == str(victim["consumo"] + 7.0)
    with pytest.raises(ValueError, match="errores"):
        data_service.apply(paths, staged)

    staged2, choose = load(paths, export, policy="usar_nuevo")
    assert choose["puede_aplicar"] and choose["hallazgos"][0]["nivel"] == "warning"
    data_service.apply(paths, staged2)
    master = read_master(paths)
    row = master[(master["planta"] == victim["planta"]) & (master["sku"] == victim["sku"]) & (master["fecha"] == victim["fecha"])]
    assert float(row["consumo"].iloc[0]) == victim["consumo"] + 7.0


def test_the_input_checks_block_a_month_gap_and_nothing_new_is_not_a_load(paths, toy):
    later = make_consumption(n_months=32)
    gap = later[(later["fecha"] >= "2025-07-01") & (later["fecha"] <= "2025-07-31")]  # June is missing
    staged, review = load(paths, gap)
    codes = {f["codigo"]: f["nivel"] for f in review["hallazgos"]}
    assert codes["month_gap"] == "error" and not review["puede_aplicar"]
    with pytest.raises(ValueError):
        data_service.apply(paths, staged)

    same = toy[(toy["fecha"] >= "2025-05-01") & (toy["fecha"] <= "2025-05-31")]
    _, review = load(paths, same)
    assert not review["puede_aplicar"] and review["hallazgos"][-1]["codigo"] == "sin_cambios" and review["resumen"]["filas_nuevas"] == 0


def test_a_partial_last_month_is_a_warning_and_a_load_without_a_master_works(tmp_path, toy):
    fresh = AppPaths(tmp_path)
    partial = toy[toy["fecha"] <= "2025-05-12"]
    staged, review = load(fresh, partial)
    parcial = [f for f in review["hallazgos"] if f["codigo"] == "mes_parcial"][0]
    assert review["puede_aplicar"] and parcial["nivel"] == "warning" and "2025-05 terminan el 2025-05-12" in parcial["mensaje"]
    assert review["resumen"]["ultimo_dia_actual"] is None and review["resumen"]["filas_actuales"] == 0
    done = data_service.apply(fresh, staged)
    assert done["respaldo"] is None and len(read_master(fresh)) == len(partial)


def test_apply_refuses_when_the_master_changed_after_the_review(paths, toy):
    staged, review = load(paths, toy[(toy["fecha"] >= "2025-06-01") & (toy["fecha"] <= "2025-06-30")])
    assert review["puede_aplicar"]
    with paths.consumption_csv.open("a", newline="") as f:
        f.write("A,SKU0/01/100gsm/2100mm/1200-1450,2025-05-31,1.0\n")
    with pytest.raises(ValueError, match="cambiaron"):
        data_service.apply(paths, staged)


def test_only_the_newest_backups_are_kept(paths, toy, monkeypatch):
    monkeypatch.setattr(data_service, "KEEP_BACKUPS", 2)
    months = [("2025-06-01", "2025-06-30"), ("2025-07-01", "2025-07-31")]
    extra = make_consumption(n_months=32)
    months.append(("2025-08-01", "2025-08-31"))
    for start, end in months:
        staged, review = load(paths, extra[(extra["fecha"] >= start) & (extra["fecha"] <= end)])
        assert review["puede_aplicar"], review["hallazgos"]
        data_service.apply(paths, staged)
        time.sleep(1.1)  # the backup name has a one second resolution
    assert len(data_service.backups(paths)) == 2


# ---- through the API ---------------------------------------------------------------------------------------------------------


def test_the_api_uploads_reviews_in_the_background_applies_and_undoes(paths, toy):
    client = TestClient(create_app(paths, today=lambda: TODAY))
    assert client.get("/api/datos").json()["actual"]["ultimo_dia"] == "2025-05-31"
    june = toy[(toy["fecha"] >= "2025-06-01") & (toy["fecha"] <= "2025-06-30")]
    upload = client.post("/api/datos/subir", files=[("files", ("junio.xlsx", xlsx(june), "application/octet-stream"))])
    assert upload.status_code == 200
    staged = upload.json()["id"]
    state = wait(client, client.post("/api/datos/revisar", json={"id": staged}).json()["trabajo"])
    assert state["estado"] == "listo" and state["resultado"]["puede_aplicar"] and state["resultado"]["resumen"]["meses_nuevos"] == ["2025-06"]
    applied = client.post("/api/datos/aplicar", json={"id": staged})
    assert applied.status_code == 200 and applied.json()["ultimo_dia"] == "2025-06-30"
    now = client.get("/api/datos").json()
    assert now["actual"]["ultimo_dia"] == "2025-06-30" and len(now["respaldos"]) == 1
    assert client.get("/api/consumo").json()["etiquetas"][-1] == "2025-06"  # the screens see the new data
    assert client.post("/api/datos/deshacer").status_code == 200
    assert client.get("/api/datos").json()["actual"]["ultimo_dia"] == "2025-05-31"
    assert client.post("/api/datos/deshacer").status_code == 404


def test_the_api_refuses_wrong_files_and_unknown_uploads(paths, toy):
    client = TestClient(create_app(paths, today=lambda: TODAY))
    bad = client.post("/api/datos/subir", files=[("files", ("datos.csv", b"x", "text/csv"))])
    assert bad.status_code == 422 and "Solo se aceptan" in bad.json()["detalle"]
    assert client.post("/api/datos/cancelar", json={"id": "nada"}).status_code == 404
    staged = client.post("/api/datos/subir", files=[("files", ("roto.xlsx", b"no", "application/octet-stream"))]).json()["id"]
    state = wait(client, client.post("/api/datos/revisar", json={"id": staged}).json()["trabajo"])
    assert state["estado"] == "error" and "No se pudo leer" in state["error"]
    assert client.post("/api/datos/aplicar", json={"id": staged}).status_code == 422
    assert client.post("/api/datos/cancelar", json={"id": staged}).status_code == 200


def test_the_messages_of_the_checks_are_in_spanish(paths):
    later = make_consumption(n_months=32)
    gap = later[(later["fecha"] >= "2025-07-01") & (later["fecha"] <= "2025-07-31")]
    _, review = load(paths, gap)
    message = {f["codigo"]: f["mensaje"] for f in review["hallazgos"]}["month_gap"]
    assert message == "Hay meses sin ninguna fila entre el primer y el último día: 2025-06."
    assert all(not any(word in f["mensaje"] for word in ("rows", "missing", "months")) for f in review["hallazgos"])


def test_the_same_day_twice_in_the_new_files_with_different_values_is_a_conflict(paths, toy):
    june = toy[(toy["fecha"] >= "2025-06-01") & (toy["fecha"] <= "2025-06-30")]
    other = june.head(1).assign(consumo=june.head(1)["consumo"] + 3.0)
    staged = data_service.stage(paths, [("a.xlsx", xlsx(june)), ("b.xlsx", xlsx(other))])["id"]
    review = data_service.review(paths, staged)
    assert review["resumen"]["conflictos"] == 1 and not review["puede_aplicar"]


def test_values_of_untouched_rows_keep_every_digit_when_the_master_is_rewritten(tmp_path):
    import numpy as np

    rng = np.random.default_rng(1)
    df = make_consumption().head(3000).assign(consumo=lambda d: rng.random(len(d)) * 100)  # 17 significant digits, the worst case for a fast parser
    master = df[df["fecha"] <= df["fecha"].quantile(0.5)]
    paths = AppPaths(tmp_path)
    write_csv(table(master), paths.consumption_csv)
    before = paths.consumption_csv.read_text().splitlines()[1:]
    new = df[df["fecha"] > master["fecha"].max()]
    staged = data_service.stage(paths, [("n.xlsx", xlsx(new))])["id"]
    data_service.review(paths, staged, "usar_nuevo")
    data_service.apply(paths, staged)
    after = set(paths.consumption_csv.read_text().splitlines()[1:])
    assert all(line in after for line in before)


def test_reading_an_export_closes_the_file_so_it_can_be_deleted_on_windows(tmp_path, toy, monkeypatch):
    from src.engine import transform_utils

    closed = []
    real = openpyxl.load_workbook

    def spy(*args, **kwargs):
        workbook = real(*args, **kwargs)
        original = workbook.close
        workbook.close = lambda: (closed.append(True), original())[1]
        return workbook

    monkeypatch.setattr(transform_utils.openpyxl, "load_workbook", spy)
    path = tmp_path / "e.xlsx"
    path.write_bytes(xlsx(toy.head(10)))
    assert len(transform_utils.read_export(path)) == 10 and closed == [True]
    path.write_bytes(b"roto")
    with pytest.raises(Exception):
        transform_utils.read_export(path)
