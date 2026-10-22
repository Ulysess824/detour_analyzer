"""Excel downloads, the forecast history and the about information."""

import io
from datetime import date

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

from src.app import __version__
from src.app.api import create_app
from src.app.export_service import XLSX
from src.app.paths import AppPaths
from tests.conftest import make_consumption
from tests.test_app_api import write_data
from tests.test_app_forecast import close_the_month, forecasted, workdir  # noqa: F401  (fixtures)

TODAY = date(2026, 10, 20)


def client_of(root):
    return TestClient(create_app(AppPaths(root), today=lambda: TODAY))


def book(response) -> openpyxl.Workbook:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    return openpyxl.load_workbook(io.BytesIO(response.content))


def frame(wb, sheet) -> pd.DataFrame:
    return pd.DataFrame(list(wb[sheet].values)[1:], columns=list(next(iter(wb[sheet].values))))


def test_the_forecast_workbook_has_a_summary_the_plantas_and_every_sku(workdir):
    response = client_of(workdir).get("/api/descargas/pronostico")
    assert 'filename="pronostico_2025-07.xlsx"' in response.headers["content-disposition"]
    wb = book(response)
    assert wb.sheetnames == ["Resumen", "Por planta", "Por SKU", "Cómo leerlo"]
    saved = pd.read_csv(workdir / "results" / "forecast_2025-07.csv")
    sku = frame(wb, "Por SKU")
    assert len(sku) == len(saved) and sku["Pronóstico (TO)"].sum() == pytest.approx(saved["ml_mean"].sum(), abs=1)
    assert (sku["Mínimo de los 3 modelos (TO)"] <= sku["Pronóstico (TO)"] + 1e-9).all() and (sku["Pronóstico (TO)"] <= sku["Máximo de los 3 modelos (TO)"] + 1e-9).all()
    planta = frame(wb, "Por planta")
    assert sorted(planta["Planta"]) == ["A", "B", "C"] and planta["Pronóstico (TO)"].sum() == pytest.approx(saved["ml_mean"].sum(), abs=1)
    summary = dict(zip(*frame(wb, "Resumen").T.values))
    assert summary["Mes pronosticado"] == "2025-07" and "Generado el" in summary and summary["Último mes real usado"] == "2025-06"
    ws = wb["Por SKU"]
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold and ws.auto_filter.ref


def test_the_forecast_workbook_uses_the_month_before_the_forecast_as_the_real(workdir):
    close_the_month(workdir)
    wb = book(client_of(workdir).get("/api/descargas/pronostico"))
    assert "Real 2025-06 (TO)" in frame(wb, "Por SKU").columns


def test_a_forecast_download_needs_a_saved_forecast(workdir, tmp_path):
    assert client_of(workdir).get("/api/descargas/pronostico", params={"mes": "2020-01"}).status_code == 404
    write_data(tmp_path, make_consumption())
    assert client_of(tmp_path).get("/api/descargas/pronostico").status_code == 404


def test_the_consumption_workbook_follows_the_filters(workdir):
    client = client_of(workdir)
    df = make_consumption()
    wb = book(client.get("/api/descargas/consumo", params={"planta": "A"}))
    assert wb.sheetnames == ["Total por mes", "Por serie"]
    total = frame(wb, "Total por mes")
    assert total["Consumo (TO)"].sum() == pytest.approx(df[df["planta"] == "A"]["consumo"].sum(), abs=1)
    series = frame(wb, "Por serie")
    assert set(series["Planta"]) == {"A"} and len(series) == 4
    one = book(client.get("/api/descargas/consumo", params={"planta": "A", "sku": "SKU1/01/100gsm/2100mm/1200-1450"}))
    assert one.sheetnames == ["Total por mes"]
    daily = book(client.get("/api/descargas/consumo", params={"nivel": "diario", "planta": "B", "desde": "2025-06-02", "hasta": "2025-06-08"}))
    assert daily.sheetnames == ["Diario"] and len(frame(daily, "Diario")) == 7
    assert client.get("/api/descargas/consumo", params={"planta": "ZZ"}).status_code == 404


def test_a_partial_month_is_noted_in_the_consumption_workbook(tmp_path):
    df = make_consumption()
    write_data(tmp_path, df[df["fecha"] <= pd.Timestamp("2025-05-12")])
    wb = book(client_of(tmp_path).get("/api/descargas/consumo"))
    assert any("incompleto" in str(row[0].value) for row in wb["Total por mes"].iter_rows(min_row=2) if row[0].value)


def test_the_tracking_workbook_needs_tracking_and_then_has_the_detail(workdir):
    client = client_of(workdir)
    assert client.get("/api/descargas/seguimiento").status_code == 404
    close_the_month(workdir)
    client.post("/api/seguimiento/medir", json={"mes": "2025-07"})
    wb = book(client.get("/api/descargas/seguimiento"))
    assert wb.sheetnames == ["Historial", "Comparación", "Por modelo", "Por planta", "Mayores errores"]
    history = frame(wb, "Historial")
    assert list(history["Mes"]) == ["2025-07"] and history["Origen"].iloc[0] == "Producción"
    assert frame(wb, "Por modelo")["Modelo"].iloc[0] == "ml_mean" and len(frame(wb, "Mayores errores")) <= 15


def test_the_history_lists_each_forecast_with_its_data_and_how_it_turned_out(workdir):
    client = client_of(workdir)
    rows = client.get("/api/historial").json()["pronosticos"]
    assert [r["mes"] for r in rows] == ["2025-07"]
    row = rows[0]
    saved = pd.read_csv(workdir / "results" / "forecast_2025-07.csv")
    assert row["total_TO"] == pytest.approx(saved["ml_mean"].sum(), abs=0.1) and row["series"] == len(saved)
    assert row["datos_hasta"] == "2025-06-30" and row["creado_en"] and row["medido"] is None and row["por_medir"] is False
    close_the_month(workdir)
    assert client.get("/api/historial").json()["pronosticos"][0]["por_medir"] is True
    client.post("/api/seguimiento/medir", json={"mes": "2025-07"})
    row = client.get("/api/historial").json()["pronosticos"][0]
    assert row["por_medir"] is False and set(row["medido"]) == {"sku", "planta", "total"}


def test_the_history_is_empty_without_forecasts_or_data(tmp_path):
    assert client_of(tmp_path).get("/api/historial").json() == {"pronosticos": []}


def test_about_gives_the_version_and_where_the_files_are(workdir):
    body = client_of(workdir).get("/api/acerca").json()
    assert body["version"] == __version__ and body["carpeta"] == str(workdir) and body["archivos"]["datos"].endswith("consumos_long.csv")
