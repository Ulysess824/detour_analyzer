"""
Read SAP consumption exports into (planta, sku, fecha, consumo) records.

Three layouts are supported and detected automatically (one wide, two long).

WIDE (sheet "RESULT"): a pivot table with one column per business day.
    - Row 7:  day headers per monthly block ("01.01.2026", ..., "#" as a blank spacer).
    - Row 9+: one row per (Sales Region, Customer, Material) with the consumption of each day.
    - The last rows ("Result", "Overall Result") are subtotals and are excluded.

LONG: already one row per (planta, material, day), with or without a "Sales Region" column.
    The header row is found by name, so either variant works. Days without movement have no
    row, so long exports carry no explicit zeros (the wide one does).

In both layouts "sku" is the material description (for example "K/01/200gsm/1800mm/1200-1400"),
not the numeric material code. Rows without a consumption value are dropped, and negative
values (stock corrections, not modelled here) are clipped to 0.
"""

from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

import openpyxl

Record = tuple[str, str, str, float]  # planta, sku, fecha, consumo

SHEET_NAME = "RESULT"
DATE_ROW = 7
FIRST_DATA_ROW = 9
FIRST_DATE_COL = 6
PLANTA_COL = 2
SKU_COL = 5  # material description, unique per planta
LONG_HEADER_SEARCH_ROWS = 3


def _parse_date(value: object) -> datetime | None:
    r"""Parse a day-header cell; blanks and the monthly '#' spacer give None."""
    if not isinstance(value, str) or value in ("", "#"):
        return None
    try:
        return datetime.strptime(value, "%d.%m.%Y")
    except ValueError:
        return None


def _find_long_header(ws) -> tuple[int, dict[str, int]] | None:
    r"""
    Find the header row of a long export and map the field names to 0-based columns.

    The header row holds 'Customer', 'Day', 'Material' and 'TO'. The material description is
    in the column right after 'Material' (its own header is blank).
    """
    rows = ws.iter_rows(min_row=1, max_row=LONG_HEADER_SEARCH_ROWS, values_only=True)
    for row_number, row in enumerate(rows, start=1):
        names = [v if isinstance(v, str) else None for v in row]
        if "Customer" in names and "Day" in names and "Material" in names and "TO" in names:
            return row_number, {
                "planta": names.index("Customer"),
                "fecha": names.index("Day"),
                "sku": names.index("Material") + 1,
                "consumo": names.index("TO"),
            }
    return None


def _read_long(ws, header_row: int, cols: dict[str, int], drop_zeros: bool) -> list[Record]:
    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        planta, sku, consumo = row[cols["planta"]], row[cols["sku"]], row[cols["consumo"]]
        if planta in (None, "Result") or sku is None or consumo is None:
            continue
        if drop_zeros and float(consumo) == 0.0:
            continue
        fecha = _parse_date(row[cols["fecha"]])
        if fecha is None:
            continue
        records.append((str(planta), str(sku), fecha.strftime("%Y-%m-%d"), max(float(consumo), 0.0)))
    return records


def _read_wide(ws, drop_zeros: bool) -> list[Record]:
    header = next(ws.iter_rows(min_row=DATE_ROW, max_row=DATE_ROW, values_only=True))
    date_by_col = {}
    for col, value in enumerate(header, start=1):
        if col >= FIRST_DATE_COL and _parse_date(value) is not None:
            date_by_col[col] = _parse_date(value)

    records = []
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        planta, sku = row[PLANTA_COL - 1], row[SKU_COL - 1]
        if planta in (None, "Result") or sku is None:
            continue
        for col, fecha in date_by_col.items():
            consumo = row[col - 1]
            if consumo is None or (drop_zeros and float(consumo) == 0.0):
                continue
            records.append((str(planta), str(sku), fecha.strftime("%Y-%m-%d"), max(float(consumo), 0.0)))
    return records


def read_export(path: Path, drop_zeros: bool = False) -> list[Record]:
    r"""
    Read one export, in any supported layout, sorted by (planta, sku, fecha).

    `drop_zeros` discards rows whose raw value is exactly 0 (before negatives are clipped).
    Long exports omit no-movement days while the wide one records them as 0, so this makes
    the wide export comparable when both kinds are combined.
    """
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = workbook[SHEET_NAME] if SHEET_NAME in workbook.sheetnames else workbook[workbook.sheetnames[0]]

    long_header = _find_long_header(ws)
    if long_header is not None:
        records = _read_long(ws, *long_header, drop_zeros=drop_zeros)
    else:
        records = _read_wide(ws, drop_zeros)
    return sorted(records, key=lambda r: (r[0], r[1], r[2]))


def combine_exports(parts: list[list[Record]]) -> list[Record]:
    r"""
    Merge the records of several exports.

    A (planta, sku, fecha) key that appears in more than one file is kept once when every copy
    has the same value (exports that overlap in time). Different values for the same key stop
    the run, because it is unclear which file is right.
    """
    merged: dict[tuple[str, str, str], float] = {}
    conflicts = 0
    for record in (r for part in parts for r in part):
        key, value = record[:3], record[3]
        if key in merged and abs(merged[key] - value) > 1e-9:
            conflicts += 1
        merged[key] = value
    if conflicts:
        raise SystemExit(f"{conflicts:_} (planta, sku, fecha) keys have different values in different inputs.")
    return sorted((*key, value) for key, value in merged.items())


def write_csv(records: list[Record], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["planta", "sku", "fecha", "consumo"])
        writer.writerows(records)
