"""The planner / internal model / ml_mean comparison as an Excel workbook: predictions with error columns, excluded rows and a summary by planta."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.utils.planner_utils import LOW_HISTORY, MISSING_INTERNAL, NO_MODEL, NO_SERIES, SUBGRADE

MISSING_TEXT = "faltante"
COMPARISON_SHEET = "Comparación"
COMPARISON_COLUMNS = ["SKU", "Planta", "Mes", "Real (TO)", "Planificador (TO)", "Modelo interno (TO)", "Nuestro modelo ml_mean (TO)"]
EXCLUDED_COLUMNS = ["SKU", "Planta", "Mes", "Planificador (TO)", "Motivo", "Explicación"]
# Formula columns after the seven data columns (D real, E planner, F internal, G ml_mean): H.. Q. They use plain cell references
# and English function names, as the file format stores them; Excel shows them in the user's language.
HELPER_FORMULAS = {
    "Tiene interno": "=ISNUMBER(F{r})",
    "Error planificador (TO)": "=E{r}-D{r}",
    "Error modelo interno (TO)": '=IF(ISNUMBER(F{r}),F{r}-D{r},"faltante")',
    "Error nuestro modelo (TO)": "=G{r}-D{r}",
    "Error abs. planificador (TO)": "=ABS(I{r})",
    "Error abs. modelo interno (TO)": '=IF(ISNUMBER(F{r}),ABS(F{r}-D{r}),"faltante")',
    "Error abs. nuestro modelo (TO)": "=ABS(K{r})",
    "Error % planificador": '=IF(D{r}=0,"",(E{r}-D{r})/D{r})',
    "Error % modelo interno": '=IF(ISNUMBER(F{r}),IF(D{r}=0,"",(F{r}-D{r})/D{r}),"faltante")',
    "Error % nuestro modelo": '=IF(D{r}=0,"",(G{r}-D{r})/D{r})',
}
SUMMARY_COLUMNS = [
    "Planta", "SKU en las tres fuentes", "Real (TO)", "Planificador (TO)", "Modelo interno (TO)", "Nuestro modelo (TO)",
    "WAPE planificador", "WAPE modelo interno", "WAPE nuestro modelo", "Acierto planificador", "Acierto modelo interno", "Acierto nuestro modelo",
    "Sesgo planificador", "Sesgo modelo interno", "Sesgo nuestro modelo",
]

REASON_TEXT = {
    SUBGRADE: ("Subgrado distinto", "El nombre del SKU no coincide con los datos de consumo porque cambia el subgrado ({detail}). No se cruzó: pendiente de confirmar con el planificador."),
    LOW_HISTORY: ("Menos de 3 meses de historia", "El SKU tiene solo {detail} de consumo antes del mes pronosticado y el modelo necesita al menos 3."),
    MISSING_INTERNAL: ("Modelo interno sin pronóstico", "{detail} Se muestra como faltante en la hoja de comparación y no cuenta en las métricas."),
    NO_SERIES: ("SKU sin consumo en los datos", "El SKU nunca tuvo consumo en los datos, no hay serie que pronosticar."),
    NO_MODEL: ("Nuestro modelo sin pronóstico", "El SKU existe pero no estaba activo el mes anterior, así que el modelo no lo pronosticó."),
}


DETAIL_TEXT = {
    "0 en todos los meses": "El modelo interno tiene 0 en todos los meses para este SKU; se toma como sin pronóstico.",
    "no está en el archivo del modelo interno": "El SKU no está en el archivo del modelo interno.",
}


def comparison_sheet(table: pd.DataFrame) -> pd.DataFrame:
    r"""SKU, planta, month, the real and the three forecasts (3 decimals, shown with 1); the internal model is "faltante" where it has no forecast."""
    out = pd.DataFrame(
        {
            "SKU": table["sku"], "Planta": table["planta"], "Mes": table["mes"], "Real (TO)": table["real"].round(3), "Planificador (TO)": table["planner"].round(3),
            "Modelo interno (TO)": table["internal"].round(3).astype(object).where(table["internal"].notna(), MISSING_TEXT),
            "Nuestro modelo ml_mean (TO)": table["ml_mean"].round(3),
        }
    )
    return out[COMPARISON_COLUMNS].sort_values(["Planta", "SKU"], ignore_index=True)


def excluded_sheet(excluded: pd.DataFrame) -> pd.DataFrame:
    r"""Every left-out row with the reason in plain words and an explanation."""
    rows = []
    for _, r in excluded.iterrows():
        reason, template = REASON_TEXT.get(r["reason"], (r["reason"], "{detail}"))
        detail = "" if pd.isna(r.get("detail")) else str(r["detail"])
        detail = DETAIL_TEXT.get(detail, detail).replace(" antes del origen", "")
        rows.append({"SKU": r["sku"], "Planta": r["planta"], "Mes": r["mes"], "Planificador (TO)": round(float(r["planner"]), 1), "Motivo": reason, "Explicación": template.format(detail=detail).replace(" ()", "")})
    return pd.DataFrame(rows, columns=EXCLUDED_COLUMNS).sort_values(["Motivo", "Planta", "SKU"], ignore_index=True)


def add_error_columns(frame: pd.DataFrame) -> pd.DataFrame:
    r"""`frame` (the comparison sheet) with live formula columns: is the internal forecast there, errors in TO, absolute errors and errors in percent."""
    out = frame.copy()
    for name, formula in HELPER_FORMULAS.items():
        out[name] = [formula.format(r=i + 2) for i in range(len(out))]
    return out


def summary_sheet(plantas: list[str], n_rows: int) -> pd.DataFrame:
    r"""
    One row per planta plus "Todas", with live formulas over the comparison sheet: SKUs, real and forecasts, WAPE, acierto and sesgo of
    every source, only on the rows where the internal model has a forecast (the SKUs of the three sources).
    """
    sheet = f"'{COMPARISON_SHEET}'!"

    def col(letter: str) -> str:
        return f"{sheet}${letter}$2:${letter}${n_rows + 1}"

    rows = []
    for k, planta in enumerate(["Todas", *plantas]):
        r = k + 2
        who = [] if planta == "Todas" else [f"{col('B')},$A{r}"]
        crit = ",".join([*who, f"{col('H')},TRUE"])
        total = lambda letter: f"=SUMIFS({col(letter)},{crit})"  # noqa: E731
        absolute = {"planificador": "L", "modelo interno": "M", "nuestro modelo": "N"}
        wape = {name: f"=SUMIFS({col(letter)},{crit})/$C{r}" for name, letter in absolute.items()}
        rows.append(
            {
                "Planta": planta, "SKU en las tres fuentes": f"=COUNTIFS({crit})", "Real (TO)": total("D"), "Planificador (TO)": total("E"), "Modelo interno (TO)": total("F"), "Nuestro modelo (TO)": total("G"),
                "WAPE planificador": wape["planificador"], "WAPE modelo interno": wape["modelo interno"], "WAPE nuestro modelo": wape["nuestro modelo"],
                "Acierto planificador": f"=1-G{r}", "Acierto modelo interno": f"=1-H{r}", "Acierto nuestro modelo": f"=1-I{r}",
                "Sesgo planificador": f"=D{r}/C{r}-1", "Sesgo modelo interno": f"=E{r}/C{r}-1", "Sesgo nuestro modelo": f"=F{r}/C{r}-1",
            }
        )
    return pd.DataFrame(rows, columns=SUMMARY_COLUMNS)


def _number_format(header: str) -> str | None:
    if "%" in header or header.startswith(("WAPE", "Acierto", "Sesgo")):
        return "0.0%"
    if "(TO)" in header:
        return "#,##0.0"
    return None


def write_workbook(path: str | Path, sheets: dict[str, pd.DataFrame]) -> None:
    r"""One sheet per entry, each an Excel table with a filter, a frozen header and tidy widths; strings starting with "=" are written as formulas."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, frame in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(list(frame.columns))
        for row in frame.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
        last = f"{get_column_letter(len(frame.columns))}{len(frame) + 1}"
        table = Table(displayName="T_" + "".join(c for c in name if c.isalnum()), ref=f"A1:{last}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        ws.freeze_panes = "A2"
        for j, col in enumerate(frame.columns, start=1):
            letter = get_column_letter(j)
            plain = [str(v) for v in frame[col].head(500) if not str(v).startswith("=")]
            ws.column_dimensions[letter].width = min(max(max([len(str(col))] + [len(v) for v in plain]) + 2, 12), 70)
            ws[f"{letter}1"].alignment = Alignment(wrap_text=True, vertical="center")
            number_format = _number_format(str(col))
            for cell in ws[letter][1:]:
                if number_format and (isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) or str(cell.value).startswith("=")):
                    cell.number_format = number_format
                if cell.value == MISSING_TEXT:
                    cell.font = Font(italic=True, color="808080")
                    cell.alignment = Alignment(horizontal="right")
        if len(frame):  # formulas that return "faltante" get the same grey italic as the typed ones
            ws.conditional_formatting.add(f"A2:{last}", CellIsRule(operator="equal", formula=[f'"{MISSING_TEXT}"'], font=Font(italic=True, color="808080")))
        ws.row_dimensions[1].height = 32
    wb.save(path)
