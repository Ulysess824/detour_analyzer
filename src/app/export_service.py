"""Excel files the app lets the user download: the forecast of a month, the real consumption of a selection and the tracking of a month."""

from __future__ import annotations

import io
import json

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.app import forecast_service, tracking_service
from src.app.consumption_service import ConsumptionStore
from src.app.paths import AppPaths

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
NUMBER = "#,##0.0"
PERCENT = "0.0"


def _write(sheets: dict[str, pd.DataFrame], formats: dict[str, dict[str, str]] | None = None, notes: dict[str, list[str]] | None = None) -> bytes:
    r"""A workbook with a styled header, frozen first row, filters, readable widths and number formats per column ({sheet: {column: format}})."""
    buffer = io.BytesIO()
    formats = formats or {}
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            for cell in ws[1]:
                cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), HEADER_FILL, Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = "A2"
            if len(frame):
                ws.auto_filter.ref = ws.dimensions
            for i, column in enumerate(frame.columns, start=1):
                letter = get_column_letter(i)
                longest = max([len(str(column))] + [len(f"{v:,.1f}" if isinstance(v, float) else str(v)) for v in frame[column].head(300)])
                ws.column_dimensions[letter].width = min(max(longest + 3, 10), 60)
                fmt = formats.get(name, {}).get(column)
                if fmt:
                    for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        row[0].number_format = fmt
            for text in (notes or {}).get(name, []):
                ws.append([text])
    return buffer.getvalue()


def _num(value) -> float | None:
    return None if value is None or (isinstance(value, float) and np.isnan(value)) else float(value)


def forecast_workbook(paths: AppPaths, store: ConsumptionStore, month: str | None, strategy: str | None = None) -> tuple[str, bytes]:
    r"""The saved forecast of a month: a summary, the plantas and every SKU with the three models, their range and the last real month."""
    saved = forecast_service.list_forecasts(paths)
    if not saved:
        raise forecast_service.ForecastMissing("Todavía no hay pronósticos. Genera el de este mes.")
    month = month or saved[0]["mes"]
    of_month = [f for f in saved if f["mes"] == month]
    strategy = strategy or (of_month[0]["estrategia"] if of_month else "h1")
    forecast = forecast_service.load_forecast(paths, month, strategy)
    last = forecast_service.reference_month(store, month)
    real = store.month_values(last)
    forecast["real_ultimo_TO"] = [real.get(key, 0.0) for key in zip(forecast["planta"], forecast["sku"])]
    forecast["minimo_TO"] = forecast[["lgbm", "xgb", "rf"]].min(axis=1)
    forecast["maximo_TO"] = forecast[["lgbm", "xgb", "rf"]].max(axis=1)
    forecast["cambio_pct"] = np.where(forecast["real_ultimo_TO"] > 0, (forecast["ml_mean"] / forecast["real_ultimo_TO"].where(forecast["real_ultimo_TO"] > 0) - 1) * 100, np.nan)
    detail = forecast.rename(columns={"planta": "Planta", "sku": "SKU", "ml_mean": "Pronóstico (TO)", "lgbm": "LightGBM (TO)", "xgb": "XGBoost (TO)", "rf": "Random Forest (TO)",
                                      "naive": "Repite el mes anterior (TO)", "minimo_TO": "Mínimo de los 3 modelos (TO)", "maximo_TO": "Máximo de los 3 modelos (TO)",
                                      "real_ultimo_TO": f"Real {last} (TO)", "cambio_pct": "Cambio contra el último real (%)"})
    detail = detail[["Planta", "SKU", "Pronóstico (TO)", "Mínimo de los 3 modelos (TO)", "Máximo de los 3 modelos (TO)", "LightGBM (TO)", "XGBoost (TO)", "Random Forest (TO)",
                     "Repite el mes anterior (TO)", f"Real {last} (TO)", "Cambio contra el último real (%)"]].sort_values(["Planta", "SKU"]).reset_index(drop=True)
    by_planta = forecast.groupby("planta").agg(ml=("ml_mean", "sum"), lgbm=("lgbm", "sum"), xgb=("xgb", "sum"), rf=("rf", "sum"), real=("real_ultimo_TO", "sum"), series=("sku", "size")).reset_index()
    by_planta["cambio"] = np.where(by_planta["real"] > 0, (by_planta["ml"] / by_planta["real"].where(by_planta["real"] > 0) - 1) * 100, np.nan)
    planta_sheet = by_planta.sort_values("ml", ascending=False).rename(columns={"planta": "Planta", "ml": "Pronóstico (TO)", "lgbm": "LightGBM (TO)", "xgb": "XGBoost (TO)", "rf": "Random Forest (TO)",
                                                                                "real": f"Real {last} (TO)", "series": "Series", "cambio": "Cambio contra el último real (%)"})
    manifest_file = paths.forecast_csv(month, strategy).with_suffix(".manifest.json")
    info = {"Mes pronosticado": month, "Estrategia": strategy, "Total pronosticado (TO)": round(float(forecast["ml_mean"].sum()), 1), "Series": int(len(forecast)), "Último mes real usado": last}
    if manifest_file.exists():
        manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
        info["Generado el"] = manifest["run"]["created_at"][:19].replace("T", " ")
        info["Datos hasta"] = manifest["data"]["last_date"]
        info["Archivo de parámetros"] = manifest["params"]["tuning_file"]
    summary = pd.DataFrame({"Dato": list(info), "Valor": [str(v) for v in info.values()]})
    howto = pd.DataFrame({"Cómo leer este archivo": [
        "Pronóstico = promedio de LightGBM, XGBoost y Random Forest (el modelo principal).",
        "Mínimo y máximo: el rango entre los tres modelos; si se separan mucho, el pronóstico es menos seguro. No es un intervalo estadístico.",
        "Repite el mes anterior: el modelo simple de referencia.",
        "Todo está en toneladas (TO).",
    ]})
    blob = _write(
        {"Resumen": summary, "Por planta": planta_sheet, "Por SKU": detail, "Cómo leerlo": howto},
        {"Por planta": {c: NUMBER for c in planta_sheet.columns if "(TO)" in c} | {"Cambio contra el último real (%)": PERCENT},
         "Por SKU": {c: NUMBER for c in detail.columns if "(TO)" in c} | {"Cambio contra el último real (%)": PERCENT}},
    )
    return f"pronostico_{month}.xlsx" if strategy == "h1" else f"pronostico_{month}_{strategy}.xlsx", blob


def consumption_workbook(store: ConsumptionStore, planta: str | None, sku: str | None, level: str, since: str | None, until: str | None) -> tuple[str, bytes]:
    r"""The real consumption of a selection: monthly (a total sheet and one row per series) or daily."""
    planta, sku, index = store.select(planta, sku)
    name = "_".join(part.replace("/", "-") for part in (planta, sku) if part) or "todo"
    if level == "diario":
        daily = store.daily(planta, sku, since, until)
        frame = pd.DataFrame({"Día": daily["etiquetas"], "Consumo (TO)": daily["valores"]})
        return f"consumo_diario_{name}.xlsx", _write({"Diario": frame}, {"Diario": {"Consumo (TO)": NUMBER}})
    monthly = store.monthly(planta, sku, breakdown=False)
    total = pd.DataFrame({"Mes": monthly["etiquetas"], "Consumo (TO)": monthly["valores"]})
    S = store.panel
    detail = pd.DataFrame(np.round(S.monthly[index], 3), columns=S.labels)
    detail.insert(0, "SKU", [S.keys[i][1] for i in index])
    detail.insert(0, "Planta", [S.keys[i][0] for i in index])
    detail = detail.sort_values(["Planta", "SKU"]).reset_index(drop=True)
    sheets = {"Total por mes": total}
    if len(index) > 1:
        sheets["Por serie"] = detail
    return f"consumo_mensual_{name}.xlsx", _write(sheets, {"Total por mes": {"Consumo (TO)": NUMBER}, "Por serie": {m: NUMBER for m in S.labels}},
                                                  {"Total por mes": [f"El mes {monthly['resumen']['mes_parcial']} está incompleto." ] if monthly["resumen"]["mes_parcial"] else []})


def tracking_workbook(paths: AppPaths, store: ConsumptionStore | None, month: str | None, strategy: str | None = None, origin: str | None = None) -> tuple[str, bytes]:
    r"""The tracking history (every strategy) and the detail of one month and strategy (models, plantas, alerts and the biggest misses)."""
    history = tracking_service.history(paths, store)
    rows = [{"Mes": r["mes"], "Origen": origin_name, "Estrategia": r["estrategia"], "Corte": r["corte"], "Acierto por SKU (%)": r["sku"], "Acierto por planta (%)": r["planta"],
             "Acierto del total (%)": r["total"], "Alertas": ", ".join(r["alertas"])}
            for origin_name, key in (("Prueba histórica", "backtest"), ("Producción", "produccion")) for r in history[key]]
    sheets = {"Historial": pd.DataFrame(rows).sort_values(["Mes", "Origen", "Estrategia"]).reset_index(drop=True)}
    comparison = [{"Estrategia": c["estrategia"], "Meses en producción": (c["produccion"] or {}).get("meses"), "SKU producción (%)": (c["produccion"] or {}).get("sku"),
                   "Planta producción (%)": (c["produccion"] or {}).get("planta"), "Total producción (%)": (c["produccion"] or {}).get("total"),
                   "Meses de prueba histórica": (c["prueba_historica"] or {}).get("meses"), "SKU prueba (%)": (c["prueba_historica"] or {}).get("sku"),
                   "Planta prueba (%)": (c["prueba_historica"] or {}).get("planta"), "Total prueba (%)": (c["prueba_historica"] or {}).get("total")} for c in history["comparacion"]]
    if comparison:
        sheets["Comparación"] = pd.DataFrame(comparison)
    detail = tracking_service.detail(paths, store, month, origin, strategy)
    sheets["Por modelo"] = pd.DataFrame([{"Modelo": m["modelo"], "SKU (%)": m["sku"], "Planta (%)": m["planta"], "Total (%)": m["total"], "Sesgo del total (%)": m["sesgo_total_pct"]} for m in detail["modelos"]])
    sheets["Por planta"] = pd.DataFrame([{"Planta": p["planta"], "Acierto (%)": p["acierto"], "Real (TO)": p["real_TO"], "Pronóstico (TO)": p["pronostico_TO"], "Sesgo (%)": p["sesgo_pct"]} for p in detail["plantas"]])
    if detail["errores"]:
        sheets["Mayores errores"] = pd.DataFrame([{"Planta": e["planta"], "SKU": e["sku"], "Real (TO)": e["real_TO"], "Pronóstico (TO)": e["pronostico_TO"], "Error (TO)": e["error_TO"]} for e in detail["errores"]])
    formats = {name: {c: PERCENT if "%" in c else NUMBER for c in frame.columns if "%" in c or "(TO)" in c} for name, frame in sheets.items()}
    suffix = "" if detail["estrategia"] == "h1" else f"_{detail['estrategia']}"
    return f"seguimiento_{detail['mes']}{suffix}.xlsx", _write(sheets, formats)
