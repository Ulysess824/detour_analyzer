"""
Transform the wide SAP consumption export (RESULT sheet, one column per business
day grouped in monthly blocks) into a long-format table with columns:
planta, sku, fecha, consumo.

Input layout (RESULT sheet):
    - Row 7:  day headers per monthly block ("01.01.2026", ..., "#" as a blank
              spacer column between months).
    - Row 8:  identifier headers (Sales Region, Customer, Customer name,
              Material, Material description).
    - Row 9+: one row per (Sales Region, Customer, Material) with the daily
              consumption value in each date column.
    - Last 2 rows: "Result" / "Overall Result" subtotals, excluded from output.

Rows with no consumption value for a given date are dropped (not filled with 0).
Negative values (stock return / partial-consumption corrections, since stock is
not modeled here) are clipped to 0 rather than dropped, to keep the daily series
continuous for downstream lag/rolling features.

Requires: openpyxl (pip install openpyxl)

Usage:
    python scripts/transform_consumos.py data/consumos_2026.xlsx data/consumos_long.csv
"""

from __future__ import annotations

import csv
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

SHEET_NAME = "RESULT"
DATE_ROW = 7
FIRST_DATA_ROW = 9
FIRST_DATE_COL = 6
PLANTA_COL = 2
SKU_COL = 4


def _parse_date(value: object) -> datetime | None:
    r"""Parse a day-header cell, skipping blanks and the monthly '#' spacer."""
    if not isinstance(value, str) or value in ("", "#"):
        return None
    try:
        return datetime.strptime(value, "%d.%m.%Y")
    except ValueError:
        return None


def transform(input_path: Path) -> list[tuple[str, str, str, float]]:
    r"""Melt the wide RESULT sheet into (planta, sku, fecha, consumo) records."""
    wb = openpyxl.load_workbook(input_path, read_only=True, data_only=True)
    ws = wb[SHEET_NAME]

    date_row = next(ws.iter_rows(min_row=DATE_ROW, max_row=DATE_ROW, values_only=True))
    date_by_col: dict[int, datetime] = {}
    for col_idx, value in enumerate(date_row, start=1):
        if col_idx < FIRST_DATE_COL:
            continue
        parsed = _parse_date(value)
        if parsed is not None:
            date_by_col[col_idx] = parsed

    records: list[tuple[str, str, str, float]] = []
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        planta = row[PLANTA_COL - 1]
        sku = row[SKU_COL - 1]
        if planta in (None, "Result") or sku is None:
            continue
        for col_idx, fecha in date_by_col.items():
            consumo = row[col_idx - 1]
            if consumo is None:
                continue
            consumo = max(float(consumo), 0.0)
            records.append((str(planta), str(sku), fecha.strftime("%Y-%m-%d"), consumo))

    records.sort(key=lambda r: (r[0], r[1], r[2]))
    return records


def main() -> None:
    if len(sys.argv) != 3:
        print("Usage: python scripts/transform_consumos.py <input.xlsx> <output.csv>")
        raise SystemExit(1)

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])

    records = transform(input_path)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["planta", "sku", "fecha", "consumo"])
        writer.writerows(records)

    print(f"Wrote {len(records):_} rows to {output_path}")


if __name__ == "__main__":
    main()
